"""Import and export of datasets.

Import reads CSV, TSV, XLSX or JSONL with an audio column and a text column
and turns the rows into records of a project. Export writes CSV, XLSX,
JSONL or a Hugging Face style audio folder (metadata.csv plus audio files).
Both run as background jobs and stream through the records in pages.
"""
import csv
import json
import re
import shutil
from pathlib import Path

from sat import audio, db
from sat.config import EXPORTS_DIR

PAGE = 500

AUDIO_COLUMNS = ("audio", "audio_path", "path", "file", "filename", "file_name", "wav", "audio_filepath")
TEXT_COLUMNS = ("text", "transcription", "transcript", "sentence", "original", "machine_text")
CORRECTED_COLUMNS = ("corrected_text", "correct_transcripts", "corrected", "final_text", "final", "label")


def sniff_columns(path):
    """Return the header and the first rows of a data file for the mapping UI."""
    rows = list(_iter_rows(path, limit=5))
    header = list(rows[0].keys()) if rows else []
    guess = {
        "audio": _guess(header, AUDIO_COLUMNS),
        "text": _guess(header, TEXT_COLUMNS),
        "corrected": _guess(header, CORRECTED_COLUMNS),
    }
    return {"columns": header, "preview": rows, "guess": guess}


def _guess(header, candidates):
    lower = {h.lower(): h for h in header}
    for c in candidates:
        if c in lower:
            return lower[c]
    return ""


def _iter_rows(path, limit=None):
    path = Path(path)
    ext = path.suffix.lower()
    count = 0
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True, data_only=True)
        ws = wb.active
        header = None
        for row in ws.iter_rows(values_only=True):
            if header is None:
                header = [str(c).strip() if c is not None else "" for c in row]
                continue
            item = {header[i]: ("" if v is None else str(v)) for i, v in enumerate(row) if i < len(header) and header[i]}
            if not any(item.values()):
                continue
            yield item
            count += 1
            if limit and count >= limit:
                break
        wb.close()
        return
    if ext in (".jsonl", ".ndjson", ".json"):
        with open(path, "r", encoding="utf-8-sig") as fh:
            if ext == ".json":
                data = json.load(fh)
                items = data if isinstance(data, list) else data.get("data") or data.get("rows") or []
                for item in items:
                    yield {k: ("" if v is None else str(v)) for k, v in item.items()}
                    count += 1
                    if limit and count >= limit:
                        break
                return
            for line in fh:
                line = line.strip()
                if not line:
                    continue
                item = json.loads(line)
                yield {k: ("" if v is None else str(v)) for k, v in item.items()}
                count += 1
                if limit and count >= limit:
                    break
        return
    delimiter = "\t" if ext == ".tsv" else ","
    with open(path, "r", encoding="utf-8-sig", newline="") as fh:
        if ext not in (".tsv",):
            try:
                sample = fh.read(4096)
                fh.seek(0)
                delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
            except (csv.Error, UnicodeDecodeError):
                fh.seek(0)
        reader = csv.DictReader(fh, delimiter=delimiter)
        for item in reader:
            yield {(k or "").strip(): ("" if v is None else str(v)) for k, v in item.items()}
            count += 1
            if limit and count >= limit:
                break


def run_import(ctx):
    params = ctx.params
    project = db.query_one("SELECT * FROM projects WHERE id=?", (params["project_id"],))
    if not project:
        raise RuntimeError("project not found")
    path = Path(params["path"])
    if not path.exists():
        raise RuntimeError("data file not found: %s" % path)
    mapping = params.get("mapping") or {}
    audio_col, text_col, corr_col = mapping.get("audio"), mapping.get("text"), mapping.get("corrected")
    if not audio_col:
        raise RuntimeError("audio column mapping is required")
    root = Path(project["root_path"])
    base = Path(params.get("base_dir") or root)
    mode = params.get("mode") or "add"
    mark_reviewed = bool(params.get("mark_reviewed"))
    stamp = db.now()
    added = updated = missing = skipped = 0
    total = 0
    for i, row in enumerate(_iter_rows(path)):
        if i % 25 == 0:
            ctx.check()
        total += 1
        raw = (row.get(audio_col) or "").strip().replace("\\", "/")
        if not raw:
            skipped += 1
            continue
        rel, abs_path = _resolve_audio(raw, root, base)
        text = (row.get(text_col) or "").strip() if text_col else ""
        corrected = (row.get(corr_col) or "").strip() if corr_col else None
        if corrected == "":
            corrected = None
        exists = abs_path is not None and abs_path.exists()
        if not exists:
            missing += 1
        f = db.query_one("SELECT * FROM files WHERE project_id=? AND rel_path=?", (project["id"], rel))
        if f is None:
            size = abs_path.stat().st_size if exists else 0
            mtime = abs_path.stat().st_mtime if exists else 0
            fid = db.insert(
                "files",
                {"project_id": project["id"], "rel_path": rel, "size": size, "mtime": mtime,
                 "status": "ready" if exists else "missing", "created_at": stamp, "updated_at": stamp},
            )
            f = db.query_one("SELECT * FROM files WHERE id=?", (fid,))
        rec = db.query_one("SELECT * FROM records WHERE file_id=? AND segment_index=0", (f["id"],))
        review = "reviewed" if (mark_reviewed and corrected) else "none"
        if rec is None:
            duration = None
            if exists:
                try:
                    duration = audio.probe_duration(abs_path)
                except Exception:
                    duration = None
            db.insert(
                "records",
                {
                    "project_id": project["id"], "file_id": f["id"], "segment_index": 0, "duration": duration,
                    "audio_path": str(abs_path if abs_path else root / rel),
                    "text": text, "corrected_text": corrected, "status": "done" if text else "pending",
                    "review": review, "provider": "import", "model": path.name, "created_at": stamp, "updated_at": stamp,
                    "transcribed_at": stamp if text else None,
                },
            )
            added += 1
        else:
            values = {"updated_at": stamp}
            if mode == "update":
                if text:
                    values["text"] = text
                    values["status"] = "done"
                if corrected is not None:
                    values["corrected_text"] = corrected
                    values["review"] = review if review != "none" else rec["review"]
                    values["version"] = (rec["version"] or 1) + 1
                db.update("records", rec["id"], values)
                updated += 1
            else:
                skipped += 1
        if total % 50 == 0:
            ctx.progress(done=total, message="%d rows read" % total)
    summary = "%d rows: %d added, %d updated, %d skipped, %d audio files missing" % (total, added, updated, skipped, missing)
    ctx.progress(done=total, total=total, message=summary, force=True)
    ctx.log(summary)
    return {"summary": summary, "rows": total, "added": added, "updated": updated, "skipped": skipped, "missing": missing}


def _resolve_audio(raw, root, base):
    """Turn a cell value into (rel_path inside the project root, absolute path)."""
    p = Path(raw)
    if not p.is_absolute():
        p = base / raw
    try:
        rel = p.resolve().relative_to(root.resolve()).as_posix()
        return rel, p
    except (ValueError, OSError):
        # Outside the project folder: keep the file name so the text survives.
        return Path(raw).name, p


def _safe_name(name):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("_") or "export"


def run_export(ctx):
    params = ctx.params
    project = db.query_one("SELECT * FROM projects WHERE id=?", (params["project_id"],))
    if not project:
        raise RuntimeError("project not found")
    fmt = params.get("format") or "csv"
    text_source = params.get("text_source") or "best"
    reviews = params.get("review") or []
    statuses = params.get("status") or []
    where = "r.project_id=?"
    args = [project["id"]]
    if reviews:
        where += " AND r.review IN (%s)" % ",".join("?" for _ in reviews)
        args += list(reviews)
    if statuses:
        where += " AND r.status IN (%s)" % ",".join("?" for _ in statuses)
        args += list(statuses)
    else:
        where += " AND r.status <> 'skipped'"
    if params.get("skip_empty", True):
        where += " AND (COALESCE(r.corrected_text, '') <> '' OR r.text <> '')"
    total = db.scalar("SELECT COUNT(*) FROM records r WHERE " + where, args)
    ctx.progress(total=total, done=0, message="exporting %d records" % total, force=True)
    name = _safe_name(params.get("name") or project["name"]) + "_" + db.now().replace(":", "").replace(" ", "_").replace("-", "")
    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    columns = ["id", "file", "segment", "start", "end", "duration", "text", "corrected_text", "final_text",
               "status", "review", "provider", "model", "language", "notes", "updated_at"]

    def rows():
        last = 0
        done = 0
        while True:
            page = db.query(
                "SELECT r.*, f.rel_path FROM records r JOIN files f ON f.id=r.file_id WHERE %s AND r.id>? ORDER BY r.id LIMIT %d" % (where, PAGE),
                args + [last],
            )
            if not page:
                return
            for r in page:
                last = r["id"]
                yield r
                done += 1
                if done % 100 == 0:
                    ctx.check()
                    ctx.progress(done=done)
            ctx.progress(done=done)

    def final_text(r):
        if text_source == "machine":
            return r["text"] or ""
        if text_source == "corrected":
            return r["corrected_text"] or ""
        return r["corrected_text"] if r["corrected_text"] not in (None, "") else (r["text"] or "")

    def flat(r, audio_ref):
        return {
            "id": r["id"], "file": r["rel_path"], "segment": r["segment_index"], "start": r["start_sec"], "end": r["end_sec"],
            "duration": r["duration"], "text": r["text"], "corrected_text": r["corrected_text"] or "", "final_text": final_text(r),
            "status": r["status"], "review": r["review"], "provider": r["provider"], "model": r["model"],
            "language": r["detected_language"], "notes": r["notes"], "updated_at": r["updated_at"], "audio": audio_ref,
        }

    if fmt == "audiofolder":
        out_dir = EXPORTS_DIR / name
        audio_dir = out_dir / "audio"
        audio_dir.mkdir(parents=True, exist_ok=True)
        meta_path = out_dir / "metadata.csv"
        copied = 0
        with open(meta_path, "w", encoding="utf-8", newline="") as fh:
            writer = csv.writer(fh)
            writer.writerow(["file_name", "transcription"] + [c for c in columns if c not in ("text",)])
            for r in rows():
                src = Path(r["audio_path"])
                target_rel = _audio_target(r)
                dst = audio_dir / target_rel
                try:
                    dst.parent.mkdir(parents=True, exist_ok=True)
                    if src.exists():
                        if params.get("convert_wav") or src.suffix.lower() != dst.suffix.lower():
                            audio.to_wav16k(src, dst)
                        else:
                            shutil.copy2(src, dst)
                        copied += 1
                except Exception as exc:
                    ctx.log("could not copy %s: %s" % (src, exc), level="warning")
                item = flat(r, "audio/" + target_rel)
                writer.writerow(["audio/" + target_rel, item["final_text"]] + [item[c] for c in columns if c != "text"])
        summary = "%d records, %d audio files written to %s" % (total, copied, out_dir)
        result = {"summary": summary, "path": str(out_dir), "records": total, "download": None}
        if params.get("zip"):
            zip_path = shutil.make_archive(str(out_dir), "zip", root_dir=str(out_dir))
            result["download"] = Path(zip_path).name
            result["summary"] = summary + " (zipped)"
        ctx.log(result["summary"])
        return result

    if fmt == "jsonl":
        out = EXPORTS_DIR / (name + ".jsonl")
        with open(out, "w", encoding="utf-8") as fh:
            for r in rows():
                item = flat(r, r["audio_path"])
                fh.write(json.dumps(item, ensure_ascii=False) + "\n")
    elif fmt == "xlsx":
        from openpyxl import Workbook
        out = EXPORTS_DIR / (name + ".xlsx")
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("records")
        ws.append(["audio"] + columns)
        for r in rows():
            item = flat(r, r["audio_path"])
            ws.append([item["audio"]] + [item[c] for c in columns])
        wb.save(str(out))
    else:
        out = EXPORTS_DIR / (name + ".csv")
        with open(out, "w", encoding="utf-8", newline="") as fh:
            writer = csv.writer(fh)
            writer.writerow(["audio"] + columns)
            for r in rows():
                item = flat(r, r["audio_path"])
                writer.writerow([item["audio"]] + [item[c] for c in columns])
    summary = "%d records written to %s" % (total, out.name)
    ctx.log(summary)
    return {"summary": summary, "path": str(out), "download": out.name, "records": total}


def _audio_target(r):
    rel = Path(r["rel_path"])
    if r["segment_index"] or r["start_sec"] is not None:
        return (rel.parent / ("%s_%04d.wav" % (rel.stem, r["segment_index"]))).as_posix()
    if Path(r["audio_path"]).suffix.lower() == ".wav":
        return rel.as_posix()
    return rel.with_suffix(".wav").as_posix()
